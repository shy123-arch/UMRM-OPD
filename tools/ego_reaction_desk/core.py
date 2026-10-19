"""Independent ego reaction slicing desk. Runtime data never belongs in Git."""
from contextlib import contextmanager
from datetime import datetime, timezone
from fractions import Fraction
from pathlib import Path
import hashlib
import json
import math
import os
import re
import shutil
import sqlite3
import subprocess
import uuid
import zipfile

HERE = Path(__file__).resolve().parent
DEFAULT_DATA = HERE.parents[2] / 'ego_reaction_annotation'
BIN = Path('/mnt/workspace/Wilson/miniconda3/envs/T2M-GPT/bin')


def now():
    return datetime.now(timezone.utc).isoformat(timespec='seconds')


def binary(name):
    found = shutil.which(name) or str(BIN / name)
    if not Path(found).is_file():
        raise ValueError(f'{name} unavailable')
    return found


def sha(path):
    with Path(path).open('rb') as f:
        return hashlib.file_digest(f, 'sha256').hexdigest()


def dump(path, value):
    path = Path(path)
    tmp = path.with_name(path.name + '.' + uuid.uuid4().hex + '.tmp')
    tmp.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')
    tmp.replace(path)


def probe(path):
    r = subprocess.run([binary('ffprobe'), '-v', 'error', '-select_streams', 'v:0',
                        '-show_entries', 'stream=avg_frame_rate,r_frame_rate,width,height:format=duration',
                        '-of', 'json', str(path)], capture_output=True, text=True, check=True, timeout=90)
    j = json.loads(r.stdout)
    s = j['streams'][0]
    fps = float(Fraction(s['avg_frame_rate'] or s['r_frame_rate']))
    return dict(duration=float(j['format']['duration']), fps=fps, width=s['width'], height=s['height'])


def cut(source, target, start, duration, proxy=False):
    target = Path(target)
    target.parent.mkdir(parents=True, exist_ok=True)
    tmp = target.with_name(target.stem + '.' + uuid.uuid4().hex + '.partial.mp4')
    cmd = [binary('ffmpeg'), '-nostdin', '-v', 'error', '-threads', '2', '-ss', f'{start:.6f}',
           '-i', str(source), '-t', f'{duration:.6f}', '-map', '0:v:0', '-an',
           '-vf', "scale=w='min(960,iw)':h=-2", '-c:v', 'libx264', '-threads', '2',
           '-preset', 'veryfast', '-crf', '24' if proxy else '20', '-pix_fmt', 'yuv420p',
           '-movflags', '+faststart', '-y', str(tmp)]
    try:
        subprocess.run(cmd, capture_output=True, check=True, timeout=300)
        info = probe(tmp)
        if abs(info['duration'] - duration) > max(.15, 2 / info['fps']):
            raise ValueError('cut duration validation failed')
        tmp.replace(target)
        return info
    finally:
        tmp.unlink(missing_ok=True)


class Conflict(ValueError):
    pass


class Store:
    def __init__(self, root=DEFAULT_DATA):
        self.root = Path(root).resolve()
        self.root.mkdir(parents=True, exist_ok=True)
        for d in ('media', 'clips', 'exports', 'logs'):
            (self.root / d).mkdir(exist_ok=True)
        self.path = self.root / 'ego_reaction.sqlite3'
        with self.db() as db:
            db.executescript('''
              PRAGMA journal_mode=WAL;
              CREATE TABLE IF NOT EXISTS items (
                id TEXT PRIMARY KEY, source TEXT NOT NULL, edit TEXT NOT NULL,
                version INTEGER NOT NULL DEFAULT 0, status TEXT NOT NULL DEFAULT 'pending',
                output TEXT, error TEXT NOT NULL DEFAULT '', updated_at TEXT NOT NULL);
              CREATE TABLE IF NOT EXISTS history (
                seq INTEGER PRIMARY KEY, id TEXT, version INTEGER, actor TEXT,
                action TEXT, payload TEXT, created_at TEXT);
              CREATE TABLE IF NOT EXISTS positions (actor TEXT PRIMARY KEY, item_id TEXT);
            ''')

    @contextmanager
    def db(self):
        db = sqlite3.connect(self.path, timeout=20)
        db.row_factory = sqlite3.Row
        try:
            yield db
            db.commit()
        except Exception:
            db.rollback()
            raise
        finally:
            db.close()

    @staticmethod
    def decode(r):
        if r is None:
            raise KeyError('item not found')
        d = dict(r)
        d['source'] = json.loads(d['source'])
        d['edit'] = json.loads(d['edit'])
        return d

    def get(self, ident):
        with self.db() as db:
            return self.decode(db.execute('SELECT * FROM items WHERE id=?', (ident,)).fetchone())

    def items(self):
        with self.db() as db:
            return [self.decode(r) for r in db.execute('SELECT * FROM items ORDER BY id')]

    def add(self, source):
        assert re.fullmatch(r'[A-Za-z0-9_-]+', source['id'])
        encoded = json.dumps(source, ensure_ascii=False, sort_keys=True)
        edit = dict(start=source['suggested_start'], end=source['anchor'],
                    reaction=source['reaction_original'], reaction_cn=source['reaction_cn'],
                    state='', note='', category='', boundary_ok=False, support_ok=False,
                    annotator_id='')
        with self.db() as db:
            old = db.execute('SELECT source FROM items WHERE id=?', (source['id'],)).fetchone()
            if old:
                if old['source'] != encoded:
                    raise Conflict('source changed; use a new batch')
                return False
            db.execute('INSERT INTO items(id,source,edit,updated_at) VALUES(?,?,?,?)',
                       (source['id'], encoded, json.dumps(edit, ensure_ascii=False), now()))
        return True

    def save(self, ident, payload):
        action = payload.get('action')
        if action not in ('draft', 'cut', 'reject'):
            raise ValueError('unknown action')
        actor = str(payload.get('annotator_id', '')).strip().upper()
        if not re.fullmatch('[A-Z0-9_-]{2,24}', actor):
            raise ValueError('请填写2–24位英文/数字标注者缩写')
        with self.db() as db:
            db.execute('BEGIN IMMEDIATE')
            old = self.decode(db.execute('SELECT * FROM items WHERE id=?', (ident,)).fetchone())
            if payload.get('version') != old['version']:
                raise Conflict('该条已被更新，请刷新后再保存；当前输入未覆盖')
            e = {k: str(payload.get(k, '')).strip() for k in ('reaction','reaction_cn','state','note','category')}
            if any(len(v) > 6000 for v in e.values()):
                raise ValueError('文本过长')
            if e['category'] not in ('','A','B','C','D','E','F'):
                raise ValueError('分类不合法')
            for k in ('start','end'):
                e[k] = float(payload[k])
                if not math.isfinite(e[k]):
                    raise ValueError('时间必须为有限数值')
            s = old['source']
            if not s['window_start'] <= e['start'] < e['end'] <= s['window_end'] + .001:
                raise ValueError('切片范围必须位于当前视频窗口内')
            e.update(annotator_id=actor, boundary_ok=payload.get('boundary_ok') is True,
                     support_ok=payload.get('support_ok') is True)
            if action == 'cut':
                if not 1 - 1e-6 <= e['end'] - e['start'] <= 15 + 1e-6:
                    raise ValueError('反应前片段须为1–15秒')
                if not e['reaction'] or not re.search('[\u4e00-\u9fff]', e['reaction_cn']):
                    raise ValueError('请保留反应描述及中文译文')
                if not e['boundary_ok'] or not e['support_ok']:
                    raise ValueError('请确认目标动作尚未开始、前情支持所列反应')
            if action == 'reject' and not e['note']:
                raise ValueError('请简要填写暂不采用的原因')
            version = old['version'] + 1
            status = {'draft':'draft','cut':'queued','reject':'rejected'}[action]
            body = json.dumps(e, ensure_ascii=False)
            db.execute('UPDATE items SET edit=?,version=?,status=?,output=NULL,error=?,updated_at=? WHERE id=?',
                       (body, version, status, '', now(), ident))
            db.execute('INSERT INTO history(id,version,actor,action,payload,created_at) VALUES(?,?,?,?,?,?)',
                       (ident,version,actor,action,body,now()))
            db.execute('INSERT OR REPLACE INTO positions VALUES(?,?)',(actor,ident))
        return self.get(ident)

    def process_one(self):
        with self.db() as db:
            db.execute('BEGIN IMMEDIATE')
            r = db.execute("SELECT * FROM items WHERE status='queued' ORDER BY updated_at LIMIT 1").fetchone()
            if not r:
                return False
            row = self.decode(r)
            db.execute("UPDATE items SET status='cutting' WHERE id=?", (row['id'],))
        path = self.root / 'clips' / f"{row['id']}_r{row['version']}.mp4"
        try:
            s,e = row['source'],row['edit']
            cut(s['source_path'],path,e['start'],e['end']-e['start'])
            result = dict(path=str(path), sha256=sha(path), bytes=path.stat().st_size,
                          duration=probe(path)['duration'])
            with self.db() as db:
                db.execute("UPDATE items SET status='ready',output=?,error='',updated_at=? WHERE id=? AND version=? AND status='cutting'",
                           (json.dumps(result),now(),row['id'],row['version']))
        except Exception as exc:
            with self.db() as db:
                db.execute("UPDATE items SET status='failed',error=?,updated_at=? WHERE id=? AND version=? AND status='cutting'",
                           (type(exc).__name__ + ': 切片失败，输入已保留；请重试或联系管理员',now(),row['id'],row['version']))
        return True

    def export(self):
        from openpyxl import Workbook
        rows = [r for r in self.items() if r['status']=='ready']
        if not rows:
            raise ValueError('没有已完成切片；先确认并保存至少一条')
        name = datetime.now(timezone.utc).strftime('ego_%Y%m%dT%H%M%S_') + uuid.uuid4().hex[:6]
        folder = self.root/'exports'/name
        folder.mkdir()
        wb = Workbook(); ws = wb.active; ws.title = 'reaction_clips'
        ws.append(['视频名称','开始时间(秒)','结束时间(秒)','reaction','reaction_ch','时长(秒)',
                   '标注者','备注','标注来源ID','源数据集','源视频UID','源视频开始时间(秒)',
                   '源视频结束时间(秒)','State','原始reaction','原始reason','分类'])
        records=[]; translations=[]
        with zipfile.ZipFile(folder/'videos.zip','w',compression=zipfile.ZIP_STORED) as videos:
            for r in rows:
                s,e=r['source'],r['edit']; o=json.loads(r['output']); p=Path(o['path'])
                if sha(p)!=o['sha256']:
                    raise ValueError('clip hash mismatch')
                videos.write(p,p.name)
                values=[p.name,0,o['duration'],e['reaction'],e['reaction_cn'],o['duration'],
                        e['annotator_id'],e['note'],f"ego:{r['id']}:r{r['version']}",s['dataset'],
                        s['video_uid'],e['start'],e['end'],e['state'],s['reaction_original'],s['reason_original'],e['category']]
                # Force text cells to strings: source annotations must not become spreadsheet formulas.
                ws.append(values)
                for c in ws[ws.max_row]:
                    if isinstance(c.value,str): c.data_type='s'
                translations.append(e['reaction_cn'])
                records.append(dict(id=r['id'],revision=r['version'],source=s,annotation=e,
                                    clip=dict(file=p.name,sha256=o['sha256'],duration=o['duration']),
                                    status='pending_independent_review',mocap_status='not_recorded',
                                    split='unassigned',cohort='ego_candidate_pool'))
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for col in ['A','D','E','H','N','O','P']: ws.column_dimensions[col].width=45
        wb.save(folder/'reaction_clips.xlsx')
        dump(folder/'translations_ego.json',dict(method='Source-preserving Chinese translation prefilled by Codex; saved/reviewed by desk annotator',
             input_sha256={'reaction_clips.xlsx':sha(folder/'reaction_clips.xlsx')},files={'reaction_clips.xlsx':translations}))
        dump(folder/'manifest.json',dict(schema='ego_reaction_candidates/v1',created_at=now(),
             count=len(records),dataset_identity='candidate_only_not_formal_636',
             time_contract='Workbook times are relative to exported pre-reaction clips; source absolute times are retained separately.',items=records))
        archive=self.root/'exports'/(name+'.zip')
        with zipfile.ZipFile(archive,'w',compression=zipfile.ZIP_STORED) as z:
            for p in folder.iterdir(): z.write(p,p.name)
        return dict(name=archive.name,count=len(rows),url='/exports/'+archive.name)
